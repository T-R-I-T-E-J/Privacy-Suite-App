"""
Sensi-Scan Unified Analyzer
Integrates detection, redaction, and scoring engines into a single interface.
All operations reference the centralized rule registry.
"""

import os
import shutil
from pathlib import Path
from typing import Dict, Any, Optional, List

try:
    from pdfminer.high_level import extract_text as extract_pdf
except ImportError:
    extract_pdf = None

try:
    import docx
except ImportError:
    docx = None

from .rule_registry import RULE_REGISTRY
from .detection_engine import DetectionEngine, Detection
from .redaction_engine import RedactionEngine
from .scoring_engine import ScoringEngine


class UnifiedAnalyzer:
    """
    Unified analyzer that orchestrates detection, redaction, and scoring.
    Single entry point for document analysis with rule-driven architecture.
    """
    
    def __init__(self, registry=None):
        """
        Initialize unified analyzer.
        
        Args:
            registry: Optional RuleRegistry (defaults to global RULE_REGISTRY)
        """
        self.registry = registry or RULE_REGISTRY
        self.detector = DetectionEngine(self.registry)
        self.redactor = RedactionEngine(self.registry)
        self.scorer = ScoringEngine(self.registry)
    
    def analyze_file(self, input_path: str, output_path: Optional[str] = None,
                    categories: Optional[List[str]] = None) -> Dict[str, Any]:
        """
        Analyze a file: detect sensitive data, calculate score, and optionally redact.
        
        Args:
            input_path: Path to input file
            output_path: Optional path for redacted output file
            categories: Optional list of categories to check
        
        Returns:
            Dictionary with analysis results
        """
        if not os.path.exists(input_path):
            return {"error": "File not found", "path": input_path}
        
        ext = os.path.splitext(input_path)[1].lower()
        
        # Extract text based on file type
        try:
            if ext == '.pdf':
                text = self._extract_pdf(input_path)
                redaction_supported = False
            elif ext == '.docx':
                text = self._extract_docx_text(input_path)
                redaction_supported = True
            elif ext in ['.txt', '.csv', '.log', '.md']:
                text = self._extract_text(input_path)
                redaction_supported = True
            elif ext in ['.jpg', '.jpeg', '.png', '.gif']:
                # Image analysis
                return self._analyze_image(input_path, output_path, categories)
            else:
                return {"error": f"Unsupported file type: {ext}", "path": input_path}
        
        except Exception as e:
            return {"error": f"Failed to read file: {str(e)}", "path": input_path}
        
        # Detect sensitive data
        detections = self.detector.detect_in_text(text, categories)
        
        # Calculate risk score
        score = self.scorer.calculate_score(detections)
        
        # Perform redaction if output path provided and supported
        redacted_file = None
        if output_path and redaction_supported and detections:
            try:
                if ext == '.docx':
                    redacted_file = self._redact_docx(input_path, output_path, categories)
                elif ext in ['.txt', '.csv', '.log', '.md']:
                    redacted_file = self._redact_text_file(input_path, output_path, categories)
            except Exception as e:
                print(f"Warning: Redaction failed: {e}")
        
        # Build detected_items list (legacy format)
        detected_items = []
        for category, info in score.breakdown.items():
            detected_items.append({
                "type": category,
                "count": info["count"],
                "context": info["context"],
                "weight": info["weight_per_item"]
            })
        
        return {
            "status": "success",
            "privacy_score": score.privacy_score,
            "risk_score": score.risk_score,
            "risk_level": score.risk_level,
            "total_detections": score.total_detections,
            "detected_items": detected_items,
            "redacted_file": redacted_file,
            "recommendations": self.scorer.get_risk_recommendations(score),
            "file_info": {
                "path": input_path,
                "size": os.path.getsize(input_path),
                "type": ext
            }
        }
    
    def analyze_text(self, text: str, categories: Optional[List[str]] = None) -> Dict[str, Any]:
        """
        Analyze raw text content.
        
        Args:
            text: Text content to analyze
            categories: Optional list of categories to check
        
        Returns:
            Dictionary with analysis results
        """
        # Detect
        detections = self.detector.detect_in_text(text, categories)
        
        # Score
        score = self.scorer.calculate_score(detections)
        
        # Redact
        redaction_result = self.redactor.redact_text_with_validation(text, categories)
        
        # Build detected_items
        detected_items = []
        for category, info in score.breakdown.items():
            detected_items.append({
                "type": category,
                "count": info["count"],
                "context": info["context"]
            })
        
        return {
            "status": "success",
            "privacy_score": score.privacy_score,
            "risk_score": score.risk_score,
            "risk_level": score.risk_level,
            "total_detections": score.total_detections,
            "detected_items": detected_items,
            "redacted_text": redaction_result.redacted_text,
            "recommendations": self.scorer.get_risk_recommendations(score)
        }
    
    # ═══════════════════════════════════════════════════════════════
    # EXTRACTION METHODS
    # ═══════════════════════════════════════════════════════════════
    
    def _extract_text(self, path: str) -> str:
        """Extract text from plain text file."""
        try:
            with open(path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            with open(path, 'r', encoding='latin-1') as f:
                return f.read()
    
    def _extract_pdf(self, path: str) -> str:
        """Extract text from PDF."""
        if not extract_pdf:
            raise ImportError("pdfminer.six is required for PDF support")
        return extract_pdf(path)
    
    def _extract_docx_text(self, path: str) -> str:
        """Extract text from DOCX."""
        if not docx:
            raise ImportError("python-docx is required for DOCX support")
        
        doc = docx.Document(path)
        
        # Collect all text
        text_parts = []
        
        # Body paragraphs
        for para in doc.paragraphs:
            text_parts.append(para.text)
        
        # Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        text_parts.append(para.text)
        
        # Headers/Footers
        for section in doc.sections:
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header:
                    for para in header.paragraphs:
                        text_parts.append(para.text)
            
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer:
                    for para in footer.paragraphs:
                        text_parts.append(para.text)
        
        return "\n".join(text_parts)
    
    # ═══════════════════════════════════════════════════════════════
    # REDACTION METHODS
    # ═══════════════════════════════════════════════════════════════
    
    def _redact_docx(self, input_path: str, output_path: str,
                    categories: Optional[List[str]] = None) -> str:
        """Redact sensitive data in DOCX file."""
        if not docx:
            raise ImportError("python-docx is required for DOCX redaction")
        
        doc = docx.Document(input_path)
        
        def redact_paragraphs(paragraph_list):
            """Helper to redact a list of paragraphs."""
            for para in paragraph_list:
                if para.text:
                    result = self.redactor.redact_text_with_validation(para.text, categories)
                    if result.redaction_count > 0:
                        para.text = result.redacted_text
        
        # Body paragraphs
        redact_paragraphs(doc.paragraphs)
        
        # Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    redact_paragraphs(cell.paragraphs)
        
        # Headers/Footers
        for section in doc.sections:
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header:
                    redact_paragraphs(header.paragraphs)
                    for table in header.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                redact_paragraphs(cell.paragraphs)
            
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer:
                    redact_paragraphs(footer.paragraphs)
                    for table in footer.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                redact_paragraphs(cell.paragraphs)
        
        doc.save(output_path)
        return output_path
    
    def _redact_text_file(self, input_path: str, output_path: str,
                         categories: Optional[List[str]] = None) -> str:
        """Redact sensitive data in plain text file."""
        text = self._extract_text(input_path)
        result = self.redactor.redact_text_with_validation(text, categories)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(result.redacted_text)
        
        return output_path
    
    def _analyze_image(self, input_path: str, output_path: Optional[str],
                      categories: Optional[List[str]] = None) -> Dict[str, Any]:
        """Analyze image file."""
        detections = self.detector.detect_in_image(input_path, categories)
        score = self.scorer.calculate_score(detections)
        
        # Copy image to output if provided
        if output_path:
            shutil.copy2(input_path, output_path)
        
        detected_items = []
        for category, info in score.breakdown.items():
            detected_items.append({
                "type": category,
                "count": info["count"],
                "context": "image"
            })
        
        return {
            "status": "success",
            "privacy_score": score.privacy_score,
            "risk_score": score.risk_score,
            "risk_level": score.risk_level,
            "total_detections": score.total_detections,
            "detected_items": detected_items,
            "redacted_file": output_path if output_path else None,
            "recommendations": self.scorer.get_risk_recommendations(score),
            "file_info": {
                "path": input_path,
                "size": os.path.getsize(input_path),
                "type": "image"
            }
        }


# ═══════════════════════════════════════════════════════════════
# CONVENIENCE FUNCTION (BACKWARD COMPATIBILITY)
# ═══════════════════════════════════════════════════════════════

def analyze_document(input_path: str, output_path: Optional[str] = None) -> Dict[str, Any]:
    """
    Convenience function for document analysis (backward compatible).
    
    Args:
        input_path: Path to input file
        output_path: Optional path for redacted output
    
    Returns:
        Analysis results dictionary
    """
    analyzer = UnifiedAnalyzer()
    return analyzer.analyze_file(input_path, output_path)


if __name__ == "__main__":
    import json
    
    print("═" * 70)
    print("SENSI-SCAN UNIFIED ANALYZER")
    print("═" * 70)
    
    # Create a test document
    test_content = """
    CONFIDENTIAL EMPLOYEE RECORD
    
    Name: John Doe
    Email: john.doe@company.com
    SSN: 123-45-6789
    Phone: 555-123-4567
    Credit Card: 4532-1488-0343-6467
    
    API Credentials:
    Stripe: sk_test_PLACEHOLDER
    AWS: AKIA_PLACEHOLDER
    """
    
    analyzer = UnifiedAnalyzer()
    
    print("\n" + "─" * 70)
    print("TEXT ANALYSIS")
    print("─" * 70)
    
    result = analyzer.analyze_text(test_content)
    
    print(f"\n📊 Results:")
    print(f"   Privacy Score: {result['privacy_score']}/100")
    print(f"   Risk Level: {result['risk_level']}")
    print(f"   Total Detections: {result['total_detections']}")
    
    print(f"\n🔍 Detected Items:")
    for item in result['detected_items']:
        print(f"   • {item['type']:20} : {item['count']} occurrence(s)")
    
    print(f"\n💡 Recommendations:")
    for rec in result['recommendations']:
        print(f"   {rec}")
    
    print(f"\n📝 Redacted Preview:")
    preview = result['redacted_text'][:300] + "..." if len(result['redacted_text']) > 300 else result['redacted_text']
    print(f"   {preview}")
